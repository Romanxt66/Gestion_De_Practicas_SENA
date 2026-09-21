"""Panel del administrador: usuarios, altas por correo, roles, fichas,
instructores, empresas, historial y exportaciones."""
import io
import os
import re

import openpyxl
import pytest
from werkzeug.security import check_password_hash

from app import db
from app.models.aprendiz import Aprendiz
from app.models.curso import Curso
from app.models.curso_aprendiz import CursoAprendiz
from app.models.curso_instructor import CursoInstructor
from app.models.empresa import Empresa
from app.models.evidencia import Evidencia
from app.models.historial_cambios import HistorialCambios
from app.models.instructor import Instructor
from app.models.notificacion import Notificacion
from app.models.rol import Rol
from app.models.usuario import Usuario
from app.models.usuario_rol import UsuarioRol
from app.models.verificacion_correo import VerificacionCorreo
from app.servicios import correo as serv_correo
from app.utils import directorio_evidencias
from conftest import (cliente_de, crear_curso, crear_usuario, enlace_de, texto, token,
                      unico)


def enviar(cliente, pagina, ruta, seguir=True, **campos):
    """POST con el token CSRF sacado de un formulario de `pagina`."""
    campos['csrf_token'] = token(cliente, pagina)
    return cliente.post(ruta, data=campos, follow_redirects=seguir)


def id_rol(nombre):
    return Rol.query.filter_by(nombre=nombre).first().id_rol


# ════════════════════════════════════════════════════════════════════════
# Empresas
# ════════════════════════════════════════════════════════════════════════
def test_crear_empresa(app, c_admin):
    nombre = f'Empresa {unico()}'
    r = enviar(c_admin, '/admin/empresas', '/admin/empresas/crear', nombre=nombre,
               nit='900123', direccion='Calle 1', telefono='3000000',
               contacto='rh@empresa.co', persona_contacto='Rosa')
    assert nombre in texto(r)
    with app.app_context():
        e = Empresa.query.filter_by(nombre=nombre).one()
        assert e.activa and e.nit == '900123' and e.persona_contacto == 'Rosa'


def test_crear_empresa_sin_nombre(c_admin):
    r = enviar(c_admin, '/admin/empresas', '/admin/empresas/crear', nombre='  ')
    assert 'El nombre es obligatorio' in texto(r)


def test_editar_empresa(app, c_admin):
    with app.app_context():
        e = Empresa(nombre=f'Vieja {unico()}', nit='1')
        db.session.add(e)
        db.session.commit()
        id_e = e.id_empresa
    r = enviar(c_admin, '/admin/empresas', f'/admin/empresas/{id_e}/editar',
               nombre='Nueva Razón', nit='2', activa='on')
    assert 'Empresa actualizada' in texto(r)
    with app.app_context():
        e = db.session.get(Empresa, id_e)
        assert e.nombre == 'Nueva Razón' and e.nit == '2' and e.activa is True
    enviar(c_admin, '/admin/empresas', f'/admin/empresas/{id_e}/editar', nombre='Nueva Razón')
    with app.app_context():
        assert db.session.get(Empresa, id_e).activa is False


def test_editar_empresa_sin_nombre_o_inexistente(app, c_admin):
    with app.app_context():
        e = Empresa(nombre=f'E {unico()}')
        db.session.add(e)
        db.session.commit()
        id_e = e.id_empresa
    r = enviar(c_admin, '/admin/empresas', f'/admin/empresas/{id_e}/editar', nombre='')
    assert 'El nombre es obligatorio' in texto(r)
    r = enviar(c_admin, '/admin/empresas', '/admin/empresas/999999/editar', nombre='x',
               seguir=False)
    assert r.status_code == 404


def test_instructor_no_gestiona_empresas(c_inst):
    t = token(c_inst, '/instructor/dashboard')
    assert c_inst.post('/admin/empresas/crear',
                       data={'nombre': 'X', 'csrf_token': t}).status_code == 403


# ════════════════════════════════════════════════════════════════════════
# Usuarios: alta por invitación
# ════════════════════════════════════════════════════════════════════════
def invitar(c_admin, **campos):
    campos.setdefault('nombres', 'Nombre')
    campos.setdefault('apellidos', 'Apellido')
    campos.setdefault('correo', f'{unico("i")}@x.co')
    campos.setdefault('password', 'clave123')
    return campos, enviar(c_admin, '/admin/usuarios', '/admin/usuarios/crear', **campos)


def test_crear_instructor_envia_confirmacion_y_no_crea_la_cuenta(app, c_admin, anonimo, correos):
    campos, r = invitar(c_admin, id_rol=_rol(app, 'instructor'), area_formacion='Teleinformática',
                        tipo_documento='CC', numero_documento=unico('') + '3',
                        telefono='3001112233')
    assert 'correo de confirmación' in texto(r)
    assert correos[-1]['para'] == campos['correo']
    assert 'Activa tu cuenta' in correos[-1]['asunto']
    with app.app_context():
        assert Usuario.query.filter_by(correo=campos['correo']).first() is None
        assert VerificacionCorreo.query.filter_by(correo=campos['correo']).first().origen == 'admin'
    assert 'Esperando confirmación de correo' in texto(c_admin.get('/admin/usuarios'))

    anonimo.get(enlace_de(correos), follow_redirects=True)
    with app.app_context():
        u = Usuario.query.filter_by(correo=campos['correo']).first()
        assert u.instructor.area_formacion == 'Teleinformática'
        assert u.telefono == '3001112233' and u.tipo_documento == 'CC'


def _rol(app, nombre):
    with app.app_context():
        return id_rol(nombre)


def test_crear_aprendiz_con_ficha_existente_lo_matricula_al_confirmar(app, datos, c_admin,
                                                                       anonimo, correos):
    campos, _ = invitar(c_admin, id_rol=_rol(app, 'aprendiz'), codigo_ficha=datos.ficha,
                        horas_requeridas='500', estado_practica='En proceso',
                        fecha_inicio_practica='2026-01-10', fecha_fin_practica='2026-07-10')
    anonimo.get(enlace_de(correos), follow_redirects=True)
    with app.app_context():
        ap = Usuario.query.filter_by(correo=campos['correo']).first().aprendiz
        assert [c.id_curso for c in ap.cursos] == [datos.id_curso]
        assert ap.horas_requeridas == 500
        assert str(ap.fecha_inicio_practica) == '2026-01-10'
        assert str(ap.fecha_fin_practica) == '2026-07-10'


def test_crear_aprendiz_con_ficha_inexistente_avisa_al_confirmar(app, c_admin, anonimo, correos):
    campos, _ = invitar(c_admin, id_rol=_rol(app, 'aprendiz'), codigo_ficha='ZZZ-NO-EXISTE')
    r = anonimo.get(enlace_de(correos), follow_redirects=True)
    assert 'todavía no está registrada' in texto(r)
    with app.app_context():
        ap = Usuario.query.filter_by(correo=campos['correo']).first().aprendiz
        assert ap.cursos == []


@pytest.mark.parametrize('cambios,mensaje', [
    ({'nombres': ''}, 'Nombres, apellidos y correo son obligatorios'),
    ({'password': '123'}, 'al menos 6'),
    ({'correo': 'sin-arroba'}, 'correo electrónico válido'),
    ({'tipo_documento': 'XX'}, 'Tipo de documento inválido'),
    ({'correo': 'aprendiz.propio@test.local'}, 'Ya existe un usuario con ese correo'),
    ({'id_rol': '99999'}, 'El rol seleccionado no existe'),
])
def test_validaciones_al_crear_usuario(c_admin, cambios, mensaje):
    campos, r = invitar(c_admin, **cambios)
    assert mensaje in texto(r)
    with c_admin.application.app_context():
        assert VerificacionCorreo.query.filter_by(correo=campos['correo']).first() is None


@pytest.mark.parametrize('cambios,mensaje', [
    ({}, 'código de ficha'),
    ({'codigo_ficha': 'CI0001', 'estado_practica': 'Inventado'}, 'Estado de práctica inválido'),
    ({'codigo_ficha': 'CI0001', 'horas_requeridas': '-5'}, 'no pueden ser negativas'),
    ({'codigo_ficha': 'CI0001', 'fecha_inicio_practica': '2026-07-10',
      'fecha_fin_practica': '2026-01-10'}, 'posterior a la de inicio'),
    ({'codigo_ficha': 'CI0001', 'fecha_inicio_practica': 'mañana'}, 'Formato de fecha inválido'),
])
def test_validaciones_de_aprendiz_al_crear_usuario(app, c_admin, cambios, mensaje):
    campos, r = invitar(c_admin, id_rol=_rol(app, 'aprendiz'), **cambios)
    assert mensaje in texto(r)
    with app.app_context():
        assert VerificacionCorreo.query.filter_by(correo=campos['correo']).first() is None


def test_crear_usuario_con_documento_repetido(app, c_admin):
    doc = unico('') + '4'
    with app.app_context():
        crear_usuario('aprendiz', numero_documento=doc)
    _, r = invitar(c_admin, numero_documento=doc)
    assert 'número de documento' in texto(r)


def test_si_el_correo_falla_no_se_crea_nada(app, c_admin, monkeypatch):
    monkeypatch.setattr(serv_correo, 'enviar_ahora', lambda *a, **k: False)
    campos, r = invitar(c_admin)
    assert 'No se pudo enviar el correo' in texto(r)
    with app.app_context():
        assert VerificacionCorreo.query.filter_by(correo=campos['correo']).first() is None


def test_reenviar_y_cancelar_una_invitacion(app, c_admin, anonimo, correos):
    campos, _ = invitar(c_admin)
    with app.app_context():
        id_pend = VerificacionCorreo.query.filter_by(correo=campos['correo']).one().id_verificacion
    antes = len(correos)
    enviar(c_admin, '/admin/usuarios', f'/admin/usuarios/pendientes/{id_pend}/reenviar')
    assert len(correos) == antes + 1 and correos[-1]['para'] == campos['correo']

    enlace = enlace_de(correos)
    r = enviar(c_admin, '/admin/usuarios', f'/admin/usuarios/pendientes/{id_pend}/cancelar')
    assert 'cancelada' in texto(r)
    with app.app_context():
        assert db.session.get(VerificacionCorreo, id_pend) is None
    assert 'cancelada' in texto(anonimo.get(enlace, follow_redirects=True))


def test_reenviar_o_cancelar_una_invitacion_inexistente(c_admin):
    for accion in ('reenviar', 'cancelar'):
        r = enviar(c_admin, '/admin/usuarios', f'/admin/usuarios/pendientes/999999/{accion}',
                   seguir=False)
        assert r.status_code == 404


def test_el_formulario_separa_los_campos_por_rol(c_admin):
    html = texto(c_admin.get('/admin/usuarios'))
    assert html.count('campos-rol') >= 2
    assert 'data-para="aprendiz"' in html and 'data-para="instructor"' in html


def test_solo_el_admin_crea_usuarios(c_inst, c_ap):
    for c in (c_inst, c_ap):
        r = c.post('/admin/usuarios/crear', data={'csrf_token': 'x'})
        assert r.status_code in (400, 403)


# ════════════════════════════════════════════════════════════════════════
# Usuarios: editar, bloquear, eliminar, cambiar de ficha
# ════════════════════════════════════════════════════════════════════════
def editar(c_admin, id_usuario, **campos):
    return enviar(c_admin, '/admin/usuarios', f'/admin/usuarios/{id_usuario}/editar', **campos)


def test_editar_usuario(app, c_admin):
    with app.app_context():
        u = crear_usuario('aprendiz')
    nuevo = f'{unico("e")}@x.co'
    r = editar(c_admin, u.id_usuario, nombres='Ana', apellidos='Gómez', correo=nuevo.upper(),
               telefono='3109998877', tipo_documento='TI', numero_documento=unico('') + '5')
    assert 'actualizado' in texto(r)
    with app.app_context():
        v = db.session.get(Usuario, u.id_usuario)
        assert (v.nombres, v.apellidos, v.correo) == ('Ana', 'Gómez', nuevo)
        assert v.telefono == '3109998877' and v.tipo_documento == 'TI'


def test_editar_usuario_cambia_la_clave_solo_si_se_envia(app, c_admin):
    with app.app_context():
        u = crear_usuario('aprendiz', clave='clave-vieja')
    base = dict(nombres='A', apellidos='B', correo=u.correo)
    editar(c_admin, u.id_usuario, **base)
    with app.app_context():
        assert check_password_hash(db.session.get(Usuario, u.id_usuario).password_hash, 'clave-vieja')
    editar(c_admin, u.id_usuario, password='clave-nueva', **base)
    with app.app_context():
        assert check_password_hash(db.session.get(Usuario, u.id_usuario).password_hash, 'clave-nueva')


@pytest.mark.parametrize('cambios,mensaje', [
    ({'nombres': ''}, 'obligatorios'),
    ({'correo': 'instructor@test.local'}, 'Ya existe otro usuario con ese correo'),
    ({'tipo_documento': 'XX'}, 'Tipo de documento inválido'),
    ({'password': '123'}, 'al menos 6'),
])
def test_validaciones_al_editar_usuario(app, c_admin, cambios, mensaje):
    with app.app_context():
        u = crear_usuario('aprendiz')
    campos = dict(nombres='A', apellidos='B', correo=u.correo)
    campos.update(cambios)
    assert mensaje in texto(editar(c_admin, u.id_usuario, **campos))


def test_editar_con_documento_de_otro(app, c_admin):
    doc = unico('') + '6'
    with app.app_context():
        crear_usuario('aprendiz', numero_documento=doc)
        u = crear_usuario('aprendiz')
    r = editar(c_admin, u.id_usuario, nombres='A', apellidos='B', correo=u.correo,
               numero_documento=doc)
    assert 'número de documento' in texto(r)


def test_editar_vuelve_al_listado_permitido(datos, c_admin):
    base = dict(nombres='Ivana', apellidos='Instructora', correo=datos.inst_correo)
    r = enviar(c_admin, '/admin/instructores', f'/admin/usuarios/{datos.inst_id}/editar',
               seguir=False, volver='instructores', **base)
    assert '/admin/instructores' in r.headers['Location']
    r = enviar(c_admin, '/admin/instructores', f'/admin/usuarios/{datos.inst_id}/editar',
               seguir=False, volver='https://sitio-externo.example/robo', **base)
    assert 'sitio-externo' not in r.headers['Location'] and '/admin/usuarios' in r.headers['Location']


def test_editar_usuario_inexistente(c_admin):
    assert editar(c_admin, 999999, nombres='A', apellidos='B', correo='a@b.co',
                  seguir=False).status_code == 404


def test_bloquear_y_reactivar_un_usuario(app, c_admin):
    with app.app_context():
        u = crear_usuario('aprendiz')
    r = enviar(c_admin, '/admin/usuarios', f'/admin/usuarios/{u.id_usuario}/toggle')
    assert 'desactivado' in texto(r)
    with app.app_context():
        assert db.session.get(Usuario, u.id_usuario).estado is False
    r = enviar(c_admin, '/admin/usuarios', f'/admin/usuarios/{u.id_usuario}/toggle')
    assert 'activado' in texto(r)
    with app.app_context():
        assert db.session.get(Usuario, u.id_usuario).estado is True


def test_no_permite_desactivar_la_propia_cuenta(app, datos, c_admin):
    r = enviar(c_admin, '/admin/usuarios', f'/admin/usuarios/{datos.admin_id}/toggle')
    assert 'propia cuenta' in texto(r)
    with app.app_context():
        assert db.session.get(Usuario, datos.admin_id).estado is True


def test_no_permite_desactivar_al_unico_superusuario(app, datos, c_admin):
    with app.app_context():
        otro = crear_usuario('aprendiz')
        db.session.add(UsuarioRol(id_usuario=otro.id_usuario, id_rol=id_rol('superusuario')))
        db.session.commit()
    c_otro = cliente_de(app, otro.id_usuario)
    with app.app_context():
        db.session.get(Usuario, otro.id_usuario).estado = False
        db.session.commit()
    # Con "otro" inactivo, el baseline es el único superusuario activo.
    r = enviar(c_otro, '/admin/usuarios', f'/admin/usuarios/{datos.admin_id}/toggle')
    assert 'único superusuario activo' in texto(r)
    with app.app_context():
        assert db.session.get(Usuario, datos.admin_id).estado is True


def _aprendiz_completo(datos, con_archivo=None):
    """Aprendiz desechable con matrícula, evidencia, aviso e historial."""
    from app.models.curso_aprendiz import CursoAprendiz as CA
    u = crear_usuario('aprendiz', ficha=datos.ficha)
    db.session.add(CA(id_curso=datos.id_curso, id_aprendiz=u.id_aprendiz))
    db.session.add(Evidencia(id_aprendiz=u.id_aprendiz, tipo='texto', contenido='ev',
                             estado='Entregada'))
    if con_archivo:
        db.session.add(Evidencia(id_aprendiz=u.id_aprendiz, tipo='archivo',
                                 contenido=f'uploads/evidencias/{con_archivo}',
                                 estado='Entregada'))
    db.session.add(Notificacion(id_usuario=u.id_usuario, mensaje='hola'))
    db.session.add(HistorialCambios(id_usuario=u.id_usuario, modulo='Prueba', accion='CREAR',
                                    descripcion=f'rastro {u.correo}'))
    db.session.commit()
    return u


def eliminar(c_admin, u, confirmacion=None):
    return enviar(c_admin, '/admin/usuarios', f'/admin/usuarios/{u.id_usuario}/eliminar',
                  confirmacion=u.correo if confirmacion is None else confirmacion)


def test_el_listado_ofrece_eliminar_y_resume_lo_que_se_pierde(c_admin):
    html = texto(c_admin.get('/admin/usuarios'))
    assert 'modalEliminarUsuario' in html and 'Se eliminará también' in html


def test_eliminar_aprendiz_borra_en_cascada_y_conserva_el_historial(app, datos, c_admin):
    nombre_archivo = f'{unico("borrar")}.pdf'
    with app.app_context():
        ruta = os.path.join(directorio_evidencias(), nombre_archivo)
        with open(ruta, 'wb') as f:
            f.write(b'%PDF')
        u = _aprendiz_completo(datos, con_archivo=nombre_archivo)
    r = eliminar(c_admin, u)
    assert 'eliminado definitivamente' in texto(r)
    with app.app_context():
        assert db.session.get(Usuario, u.id_usuario) is None
        assert db.session.get(Aprendiz, u.id_aprendiz) is None
        assert Evidencia.query.filter_by(id_aprendiz=u.id_aprendiz).count() == 0
        assert CursoAprendiz.query.filter_by(id_aprendiz=u.id_aprendiz).count() == 0
        assert Notificacion.query.filter_by(id_usuario=u.id_usuario).count() == 0
        assert UsuarioRol.query.filter_by(id_usuario=u.id_usuario).count() == 0
        rastro = HistorialCambios.query.filter_by(descripcion=f'rastro {u.correo}').one()
        assert rastro.id_usuario is None
        assert HistorialCambios.query.filter(
            HistorialCambios.descripcion.like(f'%{u.correo}%'),
            HistorialCambios.accion == 'ELIMINAR').first() is not None
        assert db.session.get(Curso, datos.id_curso) is not None
    assert not os.path.exists(ruta)
    # El historial sigue abriéndose y exportándose con entradas huérfanas
    assert c_admin.get('/admin/historial').status_code == 200
    assert c_admin.get('/admin/backup/exportar/historial').status_code == 200


def test_eliminar_instructor_conserva_la_ficha(app, c_admin):
    with app.app_context():
        u = crear_usuario('instructor')
        id_curso = crear_curso()
        db.session.add(CursoInstructor(id_curso=id_curso, id_instructor=u.id_instructor))
        db.session.commit()
    eliminar(c_admin, u)
    with app.app_context():
        assert db.session.get(Usuario, u.id_usuario) is None
        assert db.session.get(Instructor, u.id_instructor) is None
        assert CursoInstructor.query.filter_by(id_curso=id_curso).count() == 0
        assert db.session.get(Curso, id_curso) is not None


def test_eliminar_exige_reescribir_el_correo(app, c_admin):
    with app.app_context():
        u = crear_usuario('aprendiz')
    assert 'no coincide' in texto(eliminar(c_admin, u, confirmacion='otra-cosa@test.com'))
    with app.app_context():
        assert db.session.get(Usuario, u.id_usuario) is not None


def test_no_se_puede_eliminar_la_propia_cuenta(app, datos, c_admin):
    r = enviar(c_admin, '/admin/usuarios', f'/admin/usuarios/{datos.admin_id}/eliminar',
               confirmacion=datos.admin_correo)
    assert 'propia cuenta' in texto(r)
    with app.app_context():
        assert db.session.get(Usuario, datos.admin_id) is not None


def test_reenviar_el_borrado_no_da_404(app, c_admin):
    with app.app_context():
        u = crear_usuario('aprendiz')
    eliminar(c_admin, u)
    r = eliminar(c_admin, u)
    assert r.status_code == 200 and 'ya no existe' in texto(r)


def test_un_instructor_no_elimina_usuarios(app, c_inst):
    with app.app_context():
        u = crear_usuario('aprendiz')
    t = token(c_inst, '/instructor/dashboard')
    r = c_inst.post(f'/admin/usuarios/{u.id_usuario}/eliminar',
                    data={'csrf_token': t, 'confirmacion': u.correo})
    assert r.status_code == 403


def test_cambiar_de_ficha_a_un_aprendiz(app, datos, c_admin):
    with app.app_context():
        u = crear_usuario('aprendiz')
        db.session.add(CursoAprendiz(id_curso=datos.id_curso, id_aprendiz=u.id_aprendiz))
        db.session.commit()
    enviar(c_admin, '/admin/usuarios', f'/admin/usuarios/{u.id_usuario}/ficha',
           id_curso=datos.id_curso2)
    with app.app_context():
        ap = db.session.get(Aprendiz, u.id_aprendiz)
        assert [c.id_curso for c in ap.cursos] == [datos.id_curso2]
        assert ap.ficha == datos.ficha2
    r = enviar(c_admin, '/admin/usuarios', f'/admin/usuarios/{u.id_usuario}/ficha', id_curso='')
    assert 'retirado de su ficha' in texto(r)
    with app.app_context():
        ap = db.session.get(Aprendiz, u.id_aprendiz)
        assert ap.cursos == [] and ap.ficha is None


def test_cambiar_de_ficha_con_ficha_inexistente_no_deja_al_aprendiz_sin_ficha(app, datos, c_admin):
    with app.app_context():
        u = crear_usuario('aprendiz')
        db.session.add(CursoAprendiz(id_curso=datos.id_curso, id_aprendiz=u.id_aprendiz))
        db.session.commit()
    r = enviar(c_admin, '/admin/usuarios', f'/admin/usuarios/{u.id_usuario}/ficha', id_curso=999999)
    assert 'no existe' in texto(r)
    with app.app_context():
        assert CursoAprendiz.query.filter_by(id_aprendiz=u.id_aprendiz).count() == 1


def test_cambiar_ficha_a_quien_no_es_aprendiz(app, datos, c_admin):
    r = enviar(c_admin, '/admin/usuarios', f'/admin/usuarios/{datos.inst_id}/ficha',
               id_curso=datos.id_curso)
    assert 'no tiene perfil de aprendiz' in texto(r)


def test_cambiar_ficha_ignora_destinos_externos(app, datos, c_admin):
    with app.app_context():
        u = crear_usuario('aprendiz')
    r = enviar(c_admin, '/admin/usuarios', f'/admin/usuarios/{u.id_usuario}/ficha',
               seguir=False, id_curso='', next='//sitio-externo.example')
    assert 'sitio-externo' not in r.headers['Location']


def test_un_instructor_no_cambia_fichas_de_usuarios(app, c_inst):
    with app.app_context():
        u = crear_usuario('aprendiz')
    t = token(c_inst, '/instructor/dashboard')
    assert c_inst.post(f'/admin/usuarios/{u.id_usuario}/ficha',
                       data={'id_curso': '', 'csrf_token': t}).status_code == 403


# ════════════════════════════════════════════════════════════════════════
# Roles
# ════════════════════════════════════════════════════════════════════════
def test_asignar_rol_crea_el_perfil_de_instructor(app, c_admin):
    with app.app_context():
        u = crear_usuario('aprendiz')
        rid = id_rol('instructor')
    r = enviar(c_admin, '/admin/roles', '/admin/roles/asignar', id_usuario=u.id_usuario, id_rol=rid)
    assert 'asignado' in texto(r)
    with app.app_context():
        assert Instructor.query.filter_by(id_usuario=u.id_usuario).first() is not None
        assert UsuarioRol.query.filter_by(id_usuario=u.id_usuario, id_rol=rid).first()


def test_asignar_rol_crea_el_perfil_de_aprendiz(app, c_admin):
    with app.app_context():
        u = crear_usuario('instructor')
        rid = id_rol('aprendiz')
    enviar(c_admin, '/admin/roles', '/admin/roles/asignar', id_usuario=u.id_usuario, id_rol=rid)
    with app.app_context():
        ap = Aprendiz.query.filter_by(id_usuario=u.id_usuario).first()
        assert ap is not None and ap.horas_requeridas == 880


def test_asignar_un_rol_repetido_o_incompleto(app, datos, c_admin):
    with app.app_context():
        rid = id_rol('aprendiz')
    r = enviar(c_admin, '/admin/roles', '/admin/roles/asignar', id_usuario=datos.ap_id, id_rol=rid)
    assert 'ya tiene ese rol' in texto(r)
    r = enviar(c_admin, '/admin/roles', '/admin/roles/asignar', id_usuario=datos.ap_id)
    assert 'Datos incompletos' in texto(r)
    r = enviar(c_admin, '/admin/roles', '/admin/roles/asignar', id_usuario=999999, id_rol=rid)
    assert 'inexistente' in texto(r)


def test_quitar_un_rol(app, c_admin):
    with app.app_context():
        u = crear_usuario('aprendiz')
        rid = id_rol('aprendiz')
    r = enviar(c_admin, '/admin/roles', '/admin/roles/quitar', id_usuario=u.id_usuario, id_rol=rid)
    assert 'Rol removido' in texto(r)
    with app.app_context():
        assert UsuarioRol.query.filter_by(id_usuario=u.id_usuario, id_rol=rid).first() is None
    # quitarlo otra vez no falla
    assert enviar(c_admin, '/admin/roles', '/admin/roles/quitar',
                  id_usuario=u.id_usuario, id_rol=rid).status_code == 200


def test_no_te_puedes_quitar_el_rol_de_superusuario(app, datos, c_admin):
    with app.app_context():
        rid = id_rol('superusuario')
    r = enviar(c_admin, '/admin/roles', '/admin/roles/quitar', id_usuario=datos.admin_id, id_rol=rid)
    assert 'a ti mismo' in texto(r)
    with app.app_context():
        assert UsuarioRol.query.filter_by(id_usuario=datos.admin_id, id_rol=rid).first()


# ════════════════════════════════════════════════════════════════════════
# Fichas
# ════════════════════════════════════════════════════════════════════════
def crear_ficha(c_admin, **campos):
    return enviar(c_admin, '/admin/fichas', '/admin/fichas/crear', **campos)


def test_crear_ficha(app, c_admin):
    nombre, codigo = f'Ficha {unico()}', unico('F')
    r = crear_ficha(c_admin, nombre=nombre, ficha=codigo, fecha_inicio='2026-01-01',
                    fecha_fin='2026-12-01')
    assert 'Ficha creada correctamente' in texto(r)
    with app.app_context():
        c = Curso.query.filter_by(nombre=nombre).one()
        assert c.ficha == codigo and str(c.fecha_inicio) == '2026-01-01'


@pytest.mark.parametrize('campos,mensaje', [
    ({'nombre': ''}, 'El nombre es requerido'),
    ({'nombre': 'Análisis y Desarrollo de Software'}, 'Ya existe una ficha con ese nombre'),
    ({'nombre': 'Otra', 'ficha': 'CI0001'}, 'Ya existe una ficha con ese código'),
    ({'nombre': 'Fechas Mal', 'fecha_inicio': '2026-05-01', 'fecha_fin': '2026-01-01'},
     'anterior a la de inicio'),
    ({'nombre': 'Fecha Basura', 'fecha_inicio': 'no-es-fecha'}, 'inválido'),
])
def test_validaciones_al_crear_ficha(app, c_admin, campos, mensaje):
    assert mensaje in texto(crear_ficha(c_admin, **campos))
    with app.app_context():
        assert Curso.query.filter_by(nombre='Fecha Basura').first() is None
        assert Curso.query.filter_by(nombre='Fechas Mal').first() is None


def test_crear_ficha_matricula_a_quienes_la_esperaban(app, c_admin):
    codigo = unico('W-')
    with app.app_context():
        espera = crear_usuario('aprendiz', ficha=codigo)
    panel = texto(c_admin.get('/admin/fichas'))
    assert codigo in panel
    r = crear_ficha(c_admin, nombre=f'Ficha {codigo}', ficha=codigo)
    assert 'matricularon' in texto(r)
    with app.app_context():
        assert len(db.session.get(Aprendiz, espera.id_aprendiz).cursos) == 1


def test_matricular_pendientes_de_una_ficha_que_ya_existe(app, c_admin):
    codigo = unico('W-')
    with app.app_context():
        crear_curso(ficha=codigo)
        espera = crear_usuario('aprendiz', ficha=codigo)
    r = enviar(c_admin, '/admin/fichas', '/admin/fichas/matricular-pendientes', codigo=codigo)
    assert 'Se matricularon' in texto(r)
    with app.app_context():
        assert len(db.session.get(Aprendiz, espera.id_aprendiz).cursos) == 1
        assert Curso.query.filter_by(ficha=codigo).count() == 1
    r = enviar(c_admin, '/admin/fichas', '/admin/fichas/matricular-pendientes', codigo=codigo)
    assert 'No quedaban aprendices' in texto(r)


def test_matricular_pendientes_con_codigo_inexistente(c_admin):
    r = enviar(c_admin, '/admin/fichas', '/admin/fichas/matricular-pendientes', codigo='NOPE-000')
    assert 'No existe ninguna ficha' in texto(r)


def test_buscar_fichas(datos, c_admin):
    assert 'CI0001' in texto(c_admin.get('/admin/fichas?q=CI0001'))
    html = texto(c_admin.get('/admin/fichas?q=ci0002'))
    assert 'CI0002' in html and 'Análisis y Desarrollo de Software' not in html


def test_editar_ficha(app, c_admin):
    with app.app_context():
        id_c = crear_curso()
    nuevo = f'Editada {unico()}'
    r = enviar(c_admin, '/admin/fichas', f'/admin/fichas/{id_c}/editar', nombre=nuevo,
               ficha='ED-1', fecha_inicio='2026-02-01', fecha_fin='2026-08-01')
    assert 'Ficha actualizada' in texto(r)
    with app.app_context():
        c = db.session.get(Curso, id_c)
        assert c.nombre == nuevo and c.ficha == 'ED-1' and str(c.fecha_fin) == '2026-08-01'


@pytest.mark.parametrize('campos,mensaje', [
    ({'nombre': 'Análisis y Desarrollo de Software'}, 'Ya existe una ficha con ese nombre'),
    ({'ficha': 'CI0001'}, 'Ya existe una ficha con ese código'),
    ({'fecha_inicio': '2026-05-01', 'fecha_fin': '2026-01-01'}, 'anterior a la de inicio'),
    ({'fecha_inicio': 'xx'}, 'Formato de fecha inválido'),
])
def test_validaciones_al_editar_ficha(app, c_admin, campos, mensaje):
    with app.app_context():
        id_c = crear_curso()
    assert mensaje in texto(enviar(c_admin, '/admin/fichas', f'/admin/fichas/{id_c}/editar', **campos))


def test_editar_ficha_inexistente(c_admin):
    assert enviar(c_admin, '/admin/fichas', '/admin/fichas/999999/editar', nombre='x',
                  seguir=False).status_code == 404


def test_eliminar_ficha_vacia_limpia_sus_instructores(app, datos, c_admin):
    with app.app_context():
        id_c = crear_curso()
        db.session.add(CursoInstructor(id_curso=id_c, id_instructor=datos.id_instructor))
        db.session.commit()
    r = enviar(c_admin, '/admin/fichas', f'/admin/fichas/{id_c}/eliminar')
    assert 'eliminada correctamente' in texto(r)
    with app.app_context():
        assert db.session.get(Curso, id_c) is None
        assert CursoInstructor.query.filter_by(id_curso=id_c).count() == 0


def test_no_se_elimina_una_ficha_con_aprendices(app, datos, c_admin):
    r = enviar(c_admin, '/admin/fichas', f'/admin/fichas/{datos.id_curso}/eliminar')
    assert 'aprendices asignados' in texto(r)
    with app.app_context():
        assert db.session.get(Curso, datos.id_curso) is not None


def test_asignar_y_desasignar_instructor_de_una_ficha(app, datos, c_admin):
    with app.app_context():
        id_c = crear_curso()
    ruta = f'/admin/fichas/{id_c}'
    r = enviar(c_admin, '/admin/fichas', f'{ruta}/asignar-instructor',
               id_instructor=datos.id_instructor)
    assert 'quedó a cargo' in texto(r)
    r = enviar(c_admin, '/admin/fichas', f'{ruta}/asignar-instructor',
               id_instructor=datos.id_instructor)
    assert 'ya está asignado' in texto(r)
    with app.app_context():
        assert CursoInstructor.query.filter_by(id_curso=id_c).count() == 1
    r = enviar(c_admin, '/admin/fichas', f'{ruta}/desasignar-instructor/{datos.id_instructor}')
    assert 'ya no está a cargo' in texto(r)
    with app.app_context():
        assert CursoInstructor.query.filter_by(id_curso=id_c).count() == 0
    assert enviar(c_admin, '/admin/fichas', f'{ruta}/desasignar-instructor/{datos.id_instructor}',
                  seguir=False).status_code == 404


def test_asignar_instructor_sin_elegir_uno(app, c_admin):
    with app.app_context():
        id_c = crear_curso()
    r = enviar(c_admin, '/admin/fichas', f'/admin/fichas/{id_c}/asignar-instructor')
    assert 'Selecciona un instructor' in texto(r)


def test_asignar_y_quitar_vuelven_de_donde_se_venia(app, datos, c_admin):
    with app.app_context():
        id_c = crear_curso()
    detalle = f'/admin/fichas/{id_c}/detalle'
    r = enviar(c_admin, detalle, f'/admin/fichas/{id_c}/asignar-instructor', seguir=False,
               id_instructor=datos.id_instructor, volver='detalle')
    assert r.headers['Location'].endswith(detalle)
    r = enviar(c_admin, detalle, f'/admin/fichas/{id_c}/desasignar-instructor/{datos.id_instructor}',
               seguir=False)
    assert r.headers['Location'].endswith('/admin/fichas')


def test_detalle_de_ficha_del_admin(datos, c_admin):
    html = texto(c_admin.get(f'/admin/fichas/{datos.id_curso}/detalle'))
    assert 'Instructores a cargo' in html and 'Ivana Instructora' in html
    assert f'/admin/fichas/{datos.id_curso}/desasignar-instructor/{datos.id_instructor}' in html
    assert c_admin.get('/admin/fichas/999999/detalle').status_code == 404


def test_detalle_de_ficha_sin_instructor_lo_advierte(app, c_admin):
    with app.app_context():
        id_c = crear_curso()
    assert 'no tiene ningún instructor asignado' in texto(c_admin.get(f'/admin/fichas/{id_c}/detalle'))


def test_remover_un_aprendiz_de_su_ficha(app, c_admin):
    with app.app_context():
        id_c = crear_curso()
        u = crear_usuario('aprendiz')
        db.session.add(CursoAprendiz(id_curso=id_c, id_aprendiz=u.id_aprendiz))
        db.session.commit()
    r = enviar(c_admin, f'/admin/fichas/{id_c}/detalle',
               f'/admin/fichas/{id_c}/remover-aprendiz/{u.id_aprendiz}')
    assert 'removido de la ficha' in texto(r)
    with app.app_context():
        assert CursoAprendiz.query.filter_by(id_curso=id_c).count() == 0
    assert enviar(c_admin, f'/admin/fichas/{id_c}/detalle',
                  f'/admin/fichas/{id_c}/remover-aprendiz/{u.id_aprendiz}',
                  seguir=False).status_code == 404


# ════════════════════════════════════════════════════════════════════════
# Instructores
# ════════════════════════════════════════════════════════════════════════
def test_vincular_y_desvincular_ficha_desde_instructores(app, c_admin):
    with app.app_context():
        u = crear_usuario('instructor')
        id_c = crear_curso()
    ruta = f'/admin/instructores/{u.id_instructor}/fichas'
    html = texto(c_admin.get('/admin/instructores'))
    assert 'No tiene ninguna ficha asignada' in html and 'Vincular a una ficha' in html

    r = enviar(c_admin, '/admin/instructores', f'{ruta}/vincular', id_curso=id_c)
    assert 'a cargo de la ficha' in texto(r)
    with app.app_context():
        assert CursoInstructor.query.filter_by(id_instructor=u.id_instructor).count() == 1
    html = texto(c_admin.get('/admin/instructores'))
    assert f'/admin/fichas/{id_c}/detalle' in html

    r = enviar(c_admin, '/admin/instructores', f'{ruta}/vincular', id_curso=id_c)
    assert 'ya estaba a cargo' in texto(r)
    with app.app_context():
        assert CursoInstructor.query.filter_by(id_instructor=u.id_instructor).count() == 1

    r = enviar(c_admin, '/admin/instructores', f'{ruta}/{id_c}/desvincular')
    assert 'ya no está a cargo' in texto(r)
    r = enviar(c_admin, '/admin/instructores', f'{ruta}/{id_c}/desvincular')
    assert 'ya no estaba a cargo' in texto(r)


def test_vincular_sin_elegir_ficha_o_a_una_que_no_existe(app, c_admin):
    with app.app_context():
        u = crear_usuario('instructor')
    ruta = f'/admin/instructores/{u.id_instructor}/fichas/vincular'
    assert 'Selecciona la ficha' in texto(enviar(c_admin, '/admin/instructores', ruta))
    assert enviar(c_admin, '/admin/instructores', ruta, id_curso=999999,
                  seguir=False).status_code == 404


def test_desvincular_no_toca_aprendices_ni_evidencias(app, datos, c_admin):
    with app.app_context():
        u = crear_usuario('instructor')
        id_c = crear_curso()
        ap = crear_usuario('aprendiz')
        db.session.add(CursoInstructor(id_curso=id_c, id_instructor=u.id_instructor))
        db.session.add(CursoAprendiz(id_curso=id_c, id_aprendiz=ap.id_aprendiz))
        db.session.add(Evidencia(id_aprendiz=ap.id_aprendiz, tipo='texto', contenido='x',
                                 estado='Entregada'))
        db.session.commit()
    enviar(c_admin, '/admin/instructores', f'/admin/instructores/{u.id_instructor}/fichas/{id_c}/desvincular')
    with app.app_context():
        assert CursoAprendiz.query.filter_by(id_curso=id_c).count() == 1
        assert Evidencia.query.filter_by(id_aprendiz=ap.id_aprendiz).count() == 1


def test_la_ficha_vinculada_no_se_ofrece_de_nuevo(app, c_admin):
    with app.app_context():
        u = crear_usuario('instructor')
        vinculada, libre = crear_curso(), crear_curso()
        db.session.add(CursoInstructor(id_curso=vinculada, id_instructor=u.id_instructor))
        db.session.commit()
    html = texto(c_admin.get('/admin/instructores'))
    accion = f'/admin/instructores/{u.id_instructor}/fichas/vincular'
    desde = html.index(accion)
    opciones = re.findall(r'<option value="(\d+)"', html[desde:html.index('</form>', desde)])
    assert str(vinculada) not in opciones and str(libre) in opciones


def test_activar_y_desactivar_instructor(app, c_admin):
    with app.app_context():
        u = crear_usuario('instructor')
    enviar(c_admin, '/admin/instructores', f'/admin/instructores/{u.id_instructor}/toggle')
    with app.app_context():
        assert db.session.get(Instructor, u.id_instructor).activo is False
    enviar(c_admin, '/admin/instructores', f'/admin/instructores/{u.id_instructor}/toggle')
    with app.app_context():
        assert db.session.get(Instructor, u.id_instructor).activo is True


def test_editar_area_del_instructor(app, c_admin):
    with app.app_context():
        u = crear_usuario('instructor')
    r = enviar(c_admin, '/admin/instructores', f'/admin/instructores/{u.id_instructor}/area',
               area_formacion='Agroindustria')
    assert 'Área de formación actualizada' in texto(r)
    with app.app_context():
        assert db.session.get(Instructor, u.id_instructor).area_formacion == 'Agroindustria'
    enviar(c_admin, '/admin/instructores', f'/admin/instructores/{u.id_instructor}/area',
           area_formacion='  ')
    with app.app_context():
        assert db.session.get(Instructor, u.id_instructor).area_formacion is None


def test_un_instructor_no_se_vincula_fichas_solo(app, datos, c_inst):
    t = token(c_inst, '/instructor/dashboard')
    r = c_inst.post(f'/admin/instructores/{datos.id_instructor}/fichas/vincular',
                    data={'id_curso': datos.id_curso2, 'csrf_token': t})
    assert r.status_code == 403


# ════════════════════════════════════════════════════════════════════════
# Historial, dashboard y exportaciones
# ════════════════════════════════════════════════════════════════════════
def test_las_acciones_quedan_en_el_historial(app, c_admin):
    nombre = f'Empresa {unico()}'
    enviar(c_admin, '/admin/empresas', '/admin/empresas/crear', nombre=nombre)
    with app.app_context():
        h = HistorialCambios.query.filter_by(modulo='Empresas', accion='CREAR',
                                             descripcion=f'Empresa {nombre} creada').one()
        assert h.usuario is not None
    assert nombre in texto(c_admin.get('/admin/historial'))


def test_el_historial_pagina(c_admin):
    assert c_admin.get('/admin/historial?page=2').status_code == 200
    assert c_admin.get('/admin/historial?page=999').status_code == 200


def test_dashboard_del_admin(c_admin):
    r = c_admin.get('/admin/dashboard')
    assert r.status_code == 200 and 'Administrador' in texto(r)


@pytest.mark.parametrize('tipo,encabezado', [
    ('aprendices', 'Estado Práctica'), ('instructores', 'Área Formación'),
    ('empresas', 'NIT'), ('historial', 'Módulo')])
def test_exportar_a_excel(c_admin, tipo, encabezado):
    r = c_admin.get(f'/admin/backup/exportar/{tipo}')
    assert r.status_code == 200
    assert 'spreadsheetml' in r.headers['Content-Type']
    assert f'{tipo}.xlsx' in r.headers['Content-Disposition']
    hoja = openpyxl.load_workbook(io.BytesIO(r.data)).active
    assert encabezado in [c.value for c in hoja[1]]


def test_exportar_aprendices_incluye_a_los_de_la_base(c_admin, datos):
    hoja = openpyxl.load_workbook(io.BytesIO(c_admin.get('/admin/backup/exportar/aprendices').data)).active
    correos = [fila[3] for fila in hoja.iter_rows(min_row=2, values_only=True)]
    assert datos.ap_correo in correos


def test_exportar_tipo_invalido(c_admin):
    r = c_admin.get('/admin/backup/exportar/secretos', follow_redirects=True)
    assert 'Tipo de exportación no válido' in texto(r)


def test_la_exportacion_queda_registrada(app, c_admin):
    c_admin.get('/admin/backup/exportar/empresas')
    with app.app_context():
        assert HistorialCambios.query.filter_by(
            modulo='Backup', descripcion='Exportación Excel: empresas').count() >= 1
